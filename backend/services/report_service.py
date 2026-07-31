import os
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from sqlalchemy import func
from flask import current_app
from database.db import db
from models import Order, OrderItem, Product, Customer, Payment, Inventory

PERIOD_DAYS = {"daily": 1, "weekly": 7, "monthly": 30}


def _period_start(period):
    days = PERIOD_DAYS.get(period, 30)
    return datetime.now(timezone.utc) - timedelta(days=days)


def sales_report(period="monthly"):
    start = _period_start(period)
    orders = (
        Order.query.filter(Order.placed_at >= start, Order.status != "cancelled").all()
    )
    return {
        "period": period,
        "total_orders": len(orders),
        "total_revenue": float(sum((o.total_amount for o in orders), Decimal("0"))),
        "orders": [o.to_dict(include_items=False) for o in orders],
    }


def revenue_report(period="monthly"):
    start = _period_start(period)
    total = (
        db.session.query(func.coalesce(func.sum(Payment.amount), 0))
        .filter(Payment.status == "success", Payment.created_at >= start)
        .scalar()
    )
    by_method = (
        db.session.query(Payment.method, func.coalesce(func.sum(Payment.amount), 0))
        .filter(Payment.status == "success", Payment.created_at >= start)
        .group_by(Payment.method)
        .all()
    )
    return {
        "period": period,
        "total_revenue": float(total or 0),
        "by_method": {method: float(amount) for method, amount in by_method},
    }


def inventory_report():
    items = Inventory.query.all()
    return {
        "total_products": len(items),
        "low_stock_count": sum(1 for i in items if i.is_low_stock),
        "items": [i.to_dict() for i in items],
    }


def customer_report():
    customers = Customer.query.all()
    result = []
    for c in customers:
        total_spent = sum(
            (o.total_amount for o in c.orders if o.status != "cancelled"), Decimal("0")
        )
        result.append(
            {
                "customer_id": c.id,
                "name": c.user.name if c.user else None,
                "total_orders": len(c.orders),
                "total_spent": float(total_spent),
                "loyalty_points": c.loyalty_points,
            }
        )
    return {"total_customers": len(customers), "customers": result}


def top_products_report(limit=10, period="monthly"):
    start = _period_start(period)
    rows = (
        db.session.query(
            Product.id,
            Product.name,
            func.coalesce(func.sum(OrderItem.quantity), 0).label("units_sold"),
            func.coalesce(func.sum(OrderItem.subtotal), 0).label("revenue"),
        )
        .join(OrderItem, OrderItem.product_id == Product.id)
        .join(Order, Order.id == OrderItem.order_id)
        .filter(Order.placed_at >= start, Order.status != "cancelled")
        .group_by(Product.id, Product.name)
        .order_by(func.sum(OrderItem.quantity).desc())
        .limit(limit)
        .all()
    )
    return {
        "period": period,
        "top_products": [
            {"product_id": r.id, "name": r.name, "units_sold": int(r.units_sold), "revenue": float(r.revenue)}
            for r in rows
        ],
    }


def _ensure_report_folder():
    os.makedirs(current_app.config["REPORT_FOLDER"], exist_ok=True)


def export_report_pdf(title, rows, columns, filename_prefix):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    _ensure_report_folder()
    path = os.path.join(current_app.config["REPORT_FOLDER"], f"{filename_prefix}.pdf")
    doc = SimpleDocTemplate(path, pagesize=A4)
    styles = getSampleStyleSheet()
    elements = [Paragraph(title, styles["Title"]), Spacer(1, 12)]

    data = [columns] + [[str(row.get(c, "")) for c in columns] for row in rows]
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2E7D32")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F8E9")]),
            ]
        )
    )
    elements.append(table)
    doc.build(elements)
    return path


def export_report_excel(title, rows, columns, filename_prefix):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    _ensure_report_folder()
    path = os.path.join(current_app.config["REPORT_FOLDER"], f"{filename_prefix}.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]

    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2E7D32")

    for row in rows:
        ws.append([row.get(c, "") for c in columns])

    for col_cells in ws.columns:
        max_len = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
        ws.column_dimensions[col_cells[0].column_letter].width = max_len + 4

    wb.save(path)
    return path


def generate_invoice_pdf(order):
    """Generates a simple invoice PDF for a single order and returns the file path."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    _ensure_report_folder()
    path = os.path.join(current_app.config["REPORT_FOLDER"], f"invoice_order_{order.id}.pdf")
    doc = SimpleDocTemplate(path, pagesize=A4)
    styles = getSampleStyleSheet()
    elements = [
        Paragraph("FreshMart - Tax Invoice", styles["Title"]),
        Paragraph(f"Order #{order.id}  |  Status: {order.status}", styles["Normal"]),
        Paragraph(f"Delivery Address: {order.delivery_address}", styles["Normal"]),
        Spacer(1, 16),
    ]

    data = [["Product", "Qty", "Unit Price", "Subtotal"]]
    for item in order.items:
        data.append(
            [item.product.name if item.product else "", item.quantity,
             f"₹{item.unit_price}", f"₹{item.subtotal}"]
        )
    data.append(["", "", "Total", f"₹{order.total_amount}"])

    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2E7D32")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
            ]
        )
    )
    elements.append(table)
    doc.build(elements)
    return path
